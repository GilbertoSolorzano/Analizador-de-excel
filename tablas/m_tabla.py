import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
_table_counter = [0]  

def _make_table(ws, df, startrow, suffix):
        if df.empty:
            return
        _table_counter[0] += 1
        nrows, ncols = df.shape
        ref = f"A{startrow + 1}:{get_column_letter(ncols)}{startrow + nrows + 1}"
        table_name = f"Tabla_{_table_counter[0]}_{suffix}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(tbl)

def autofit_columns(ws):
    for column in ws.columns:
        max_length = 0
        column_letter = None
        for cell in column:
            if column_letter is None:
                column_letter = cell.column_letter
            try:
                if cell.value:
                    cell_length = len(str(cell.value))
                    if cell_length > max_length:
                        max_length = cell_length
            except:
                pass
        if column_letter:
            ws.column_dimensions[column_letter].width = max_length + 2